"""Offline tests for the draft Course Director applicant view; no REDCap or SMTP calls."""
import importlib.util
import io
import os
import unittest
from pathlib import Path
from unittest.mock import patch, Mock

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "process_project79_applicant_packs.py"
with patch.dict(os.environ, {"REDCAP_PROJECT75_TOKEN": "offline", "REDCAP_PROJECT79_TOKEN": "offline"}):
    spec = importlib.util.spec_from_file_location("applicant_pack", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)


class ApplicantPackPilotTests(unittest.TestCase):
    def test_hidden_other_education_is_not_exported(self):
        applicants = [{"record_id": "1", "applied_run_id": "206-1", "highest_education": "3"}]
        self.assertNotIn("other_education", module.director_columns(applicants))

    def test_conditional_field_visible_for_any_eligible_applicant(self):
        applicants = [
            {"record_id": "1", "applied_run_id": "206-1", "highest_education": "3"},
            {"record_id": "2", "applied_run_id": "206-1", "highest_education": "99", "other_education": "Clinical qualification"},
        ]
        self.assertIn("other_education", module.director_columns(applicants))
        workbook = load_workbook(io.BytesIO(module.build_workbook({}, applicants)))
        sheet = workbook["All Applicants"]
        headings = [c.value for c in sheet[1]]
        col = headings.index("other_education") + 1
        self.assertEqual(sheet.cell(2, col).value, None)
        self.assertEqual(sheet.cell(3, col).value, "Clinical qualification")

    def test_uploads_only_when_requested_and_present(self):
        applicant = {
            "applied_run_id": "206-1", "course_requires_cv": "1",
            "course_requires_certificate": "0", "upload_cv": "123",
            "upload_certificate": "456", "upload_support_letter": "",
        }
        self.assertEqual(module.applicable_upload_fields(applicant), ["upload_cv"])
        applicant["course_requires_cv"] = "0"
        self.assertEqual(module.applicable_upload_fields(applicant), [])
        applicant["upload_support_letter"] = "789"
        self.assertEqual(module.applicable_upload_fields(applicant), ["upload_support_letter"])

    def test_file_export_is_read_only_and_sanitizes_filename(self):
        response = Mock()
        response.content = b"%PDF-synthetic-test-data"
        response.headers = {
            "Content-Type": "application/pdf",
            "Content-Disposition": 'attachment; filename="../../unsafe.pdf"',
        }
        response.raise_for_status.return_value = None
        with patch.object(module.requests, "post", return_value=response) as post:
            content, name = module.fetch_private_upload("123", "upload_cv")
        self.assertEqual(content, b"%PDF-synthetic-test-data")
        self.assertNotIn("/", name)
        self.assertEqual(post.call_args.kwargs["data"]["action"], "export")
        self.assertEqual(post.call_args.kwargs["data"]["content"], "file")
        with self.assertRaises(ValueError):
            module.fetch_private_upload("123", "certificate_file")

    def test_file_export_rejects_json_error_payload(self):
        response = Mock()
        response.content = b'{"error":"file not found"}'
        response.headers = {"Content-Type": "application/json"}
        response.raise_for_status.return_value = None
        with patch.object(module.requests, "post", return_value=response):
            with self.assertRaises(RuntimeError):
                module.fetch_private_upload("123", "upload_cv")

    def test_choice_labels_and_return_headers_unchanged(self):
        applicants = [{"record_id": "1", "applied_run_id": "206-1", "sex": "2", "full_name": "Pilot Applicant"}]
        workbook = load_workbook(io.BytesIO(module.build_workbook({}, applicants, {"sex": {"2": "Female"}})))
        sheet = workbook["All Applicants"]
        headings = [c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(2, headings.index("sex") + 1).value, "Female")
        self.assertEqual([c.value for c in workbook["Selection Return"][1]],
                         ["record_id", "full_name", "selected_for_batch", "selected_batch_id",
                          "batch_role", "selection_date", "selection_by"])
        self.assertEqual([c.value for c in workbook["Cert-Graduands Return"][1]],
                         ["record_id", "full_name", "gepg_control_no", "attendance_verified",
                          "attendance_verified_date", "certificate_approved", "certification_notes"])


if __name__ == "__main__":
    unittest.main()
