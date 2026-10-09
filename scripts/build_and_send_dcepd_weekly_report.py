#!/usr/bin/env python3
"""Build and optionally email the private DCEPD weekly management workbook."""
from __future__ import annotations

import argparse
import csv
import json
import os
import smtplib
import ssl
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TZ = ZoneInfo("Africa/Dar_es_Salaam")
API = os.getenv("REDCAP_API_URL", "https://utafiti.muhas.ac.tz/api/")


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def metadata75() -> list[dict[str, str]]:
    token = os.getenv("REDCAP_PROJECT75_TOKEN", "")
    if not token:
        return []
    body = urlencode({
        "token": token,
        "content": "metadata",
        "format": "json",
        "returnFormat": "json",
    }).encode()
    try:
        with urlopen(Request(API, data=body), timeout=90) as response:
            data = json.load(response)
        return data if isinstance(data, list) else []
    except Exception as exc:
        print(f"warning=project75_metadata_unavailable:{exc}")
        return []


def parse_choices(metadata: list[dict[str, str]], field_name: str) -> dict[str, str]:
    item = next((x for x in metadata if x.get("field_name") == field_name), None)
    if not item:
        return {}
    out = {}
    for part in (item.get("select_choices_or_calculations") or "").split("|"):
        if "," not in part:
            continue
        value, label = part.split(",", 1)
        out[value.strip()] = label.strip()
    return out


def add_sheet(wb: Workbook, title: str, rows: list[dict[str, str]], preferred: list[tuple[str, str]] | None = None):
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["No items"])
        return ws

    if preferred:
        fields = [x[0] for x in preferred]
        labels = [x[1] for x in preferred]
    else:
        fields = list(rows[0].keys())
        labels = fields

    ws.append(labels)
    for row in rows:
        ws.append([row.get(f, "") for f in fields])

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col):
            max_len = max(max_len, len(str(row[0].value or "")))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 48)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def build_workbook(management_dir: Path, qc_path: Path | None = None) -> Path:
    summary_path = management_dir / "management_summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    actions = read_csv(management_dir / "action_list.csv")
    overdue = read_csv(management_dir / "overdue_reaccreditation.csv")
    qc = read_csv(qc_path) if qc_path and qc_path.exists() else []

    directors = parse_choices(metadata75(), "course_director_id")
    for rows in (actions, overdue):
        for row in rows:
            row["course_director"] = directors.get(row.get("course_director_id", ""), row.get("course_director_id", ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Management Summary"
    ws.append(["DCEPD Weekly Management Report", ""])
    ws.merge_cells("A1:B1")
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    summary_rows = [
        ("Generated", datetime.now(TZ).strftime("%d %B %Y, %H:%M EAT")),
        ("Verified source run", summary.get("source_run_time_eat", "")),
        ("Automation status", summary.get("automation_status", "")),
        ("Public catalogue courses", summary.get("public_catalogue_courses", "")),
        ("QC errors", summary.get("qc_errors", "")),
        ("QC warnings", summary.get("qc_warnings", "")),
        ("QC information items", summary.get("qc_info", "")),
        ("Reaccreditation overdue", summary.get("reaccreditation_overdue", "")),
        ("Overdue with newer submission", summary.get("reaccreditation_overdue_in_progress", "")),
        ("Overdue with no newer submission", summary.get("reaccreditation_overdue_no_new_submission", "")),
        ("Reaccreditation due soon", summary.get("reaccreditation_due_soon", "")),
        ("Other action-list items", summary.get("action_list_rows", "")),
        ("Duplicate run-date reviews", summary.get("duplicate_run_date_reviews", 0)),
        ("Project 79 course-choice QC items", len(qc)),
    ]
    for a, b in summary_rows:
        ws.append([a, b])
    for cell in ws["A"][1:]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 75
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    preferred = [
        ("record_id", "Record ID"),
        ("course_code", "Course code"),
        ("course_name", "Course name"),
        ("course_director", "Course Director"),
        ("contact_email", "Director email"),
        ("contact_phone", "Director phone"),
        ("severity", "Severity"),
        ("issue_code", "Issue"),
        ("detail", "Action / detail"),
        ("accreditation_date", "Accreditation date"),
        ("date_submitted", "Date submitted"),
        ("reaccreditation_progress", "Re-accreditation progress"),
        ("last_date_conducted", "Last run"),
        ("interested_applicants", "Interested applicants"),
        ("public_catalogue", "Public catalogue"),
    ]
    add_sheet(wb, "Action List", actions, preferred)
    add_sheet(wb, "Overdue Reaccreditation", overdue, preferred)

    if qc:
        add_sheet(
            wb,
            "Project79 Choice QC",
            qc,
            [
                ("issue_type", "Issue type"),
                ("p79_choice_value", "P79 choice"),
                ("p75_record_id", "P75 record"),
                ("current_p79_label", "Current P79 label"),
                ("expected_p79_label", "Expected label"),
                ("action", "Action"),
            ],
        )

    date_text = datetime.now(TZ).strftime("%Y-%m-%d")
    out = management_dir / f"DCEPD_Weekly_Management_Report_{date_text}.xlsx"
    wb.save(out)
    print(f"workbook={out}")
    return out


def send_email(workbook: Path, summary: dict):
    recipients = []
    for key in ("DCEPD_EMAIL_1", "DCEPD_EMAIL_2", "DCEPD_EMAIL_3", "DCEPD_EMAIL_4"):
        value = (os.getenv(key) or "").strip()
        if value and value not in recipients:
            recipients.append(value)
    required = {
        "DCEPD_SMTP_HOST": os.getenv("DCEPD_SMTP_HOST", "").strip(),
        "DCEPD_SMTP_USERNAME": os.getenv("DCEPD_SMTP_USERNAME", "").strip(),
        "DCEPD_SMTP_PASSWORD": os.getenv("DCEPD_SMTP_PASSWORD", "").strip(),
    }
    missing = [k for k, v in required.items() if not v]
    if not recipients:
        missing.append("DCEPD_EMAIL_1..4")
    if missing:
        raise RuntimeError("Missing email configuration: " + ", ".join(missing))

    host = required["DCEPD_SMTP_HOST"]
    username = required["DCEPD_SMTP_USERNAME"]
    password = required["DCEPD_SMTP_PASSWORD"]
    port = int(os.getenv("DCEPD_SMTP_PORT") or ("465" if "gmail" in host.lower() else "587"))

    msg = EmailMessage()
    msg["Subject"] = "DCEPD Weekly Management Report - " + datetime.now(TZ).strftime("%d %B %Y")
    msg["From"] = username
    msg["To"] = ", ".join(recipients)
    msg.set_content(
        "Dear DCEPD team,\n\n"
        "Please find attached the weekly DCEPD management report generated from the latest verified registry refresh. "
        "The workbook includes the Management Summary, Action List, Overdue Reaccreditation and Project 79 course-choice QC where available.\n\n"
        f"Verified source: {summary.get('source_run_time_eat', '')}\n"
        f"Automation status: {summary.get('automation_status', '')}\n"
        f"Public catalogue courses: {summary.get('public_catalogue_courses', '')}\n"
        f"QC errors: {summary.get('qc_errors', '')}\n\n"
        "This is an internal management report.\n"
    )
    msg.add_attachment(
        workbook.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=workbook.name,
    )

    context = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=context, timeout=60) as smtp:
            smtp.login(username, password)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=60) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(username, password)
            smtp.send_message(msg)
    print("email_sent=true")
    print("recipients=" + ",".join(recipients))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--management-dir", default=".project75-runtime/management")
    ap.add_argument("--qc", default="project79-choice-qc/project79_course_choice_qc.csv")
    ap.add_argument("--send", action="store_true")
    args = ap.parse_args()

    management_dir = Path(args.management_dir)
    qc_path = Path(args.qc)
    summary = json.loads((management_dir / "management_summary.json").read_text(encoding="utf-8"))
    workbook = build_workbook(management_dir, qc_path)
    if args.send:
        send_email(workbook, summary)


if __name__ == "__main__":
    main()
