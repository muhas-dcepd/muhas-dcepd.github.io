#!/usr/bin/env python3
"""Generate and email per-course Project 79 applicant packs.

Read-only against REDCap:
- Project 75 supplies authoritative course/director/contact metadata.
- Project 79 supplies applicant/selection/certification data.
- This script never writes to REDCap.

State is kept in automation/applicant-pack-state.json so packs are sent only when:
1) bootstrap has not yet covered the historic Project 79 applications for that course; or
2) at least 5 current-FY applications are new since the previous pack; or
3) on Friday, 1-4 current-FY applications are new since the previous pack.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import smtplib
import ssl
from collections import defaultdict
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

REDCAP_URL = os.getenv("REDCAP_API_URL", "https://utafiti.muhas.ac.tz/api/")
P75_TOKEN = os.environ["REDCAP_PROJECT75_TOKEN"]
P79_TOKEN = os.environ["REDCAP_PROJECT79_TOKEN"]
TZ = ZoneInfo("Africa/Dar_es_Salaam")

STATE_PATH = Path(os.getenv("DCEPD_APPLICANT_PACK_STATE", "automation/applicant-pack-state.json"))
OUT_DIR = Path(os.getenv("DCEPD_APPLICANT_PACK_OUT", ".applicant-pack-output"))

P75_FIELDS = [
    "record_id", "course_code", "course_name", "course_director_id",
    "contact_email", "contact_phone", "vote_code",
]
P79_FIELDS = [
    "record_id", "applied_course_id", "application_date", "full_name", "email",
    "phone_number", "institution_name", "job_title", "payment_reference",
    "fee_verified", "selected_for_batch", "selected_batch_id", "batch_role",
    "selection_date", "selection_by", "attendance_verified",
    "attendance_verified_date", "certificate_approved", "gepg_control_no",
    "certification_notes",
]
PLACEHOLDER_DIRECTOR_TEXT = ("not chosen", "not specified", "xxx")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

IMPORT_COMMENT = (
    "For Coordinator/Admin: when the Course Director returns this workbook, save this "
    "sheet as CSV UTF-8 (Comma delimited) (*.csv) before importing into REDCap. "
    "Do not rename the REDCap field-name columns and do not delete the record_id column. "
    "Import only the intended return sheet, not the whole workbook."
)


def redcap_post(token: str, data: dict) -> requests.Response:
    payload = dict(data)
    payload["token"] = token
    response = requests.post(REDCAP_URL, data=payload, timeout=90)
    response.raise_for_status()
    return response


def export_metadata(token: str) -> list[dict]:
    return redcap_post(token, {
        "content": "metadata",
        "format": "json",
        "returnFormat": "json",
    }).json()


def export_records(token: str, fields: list[str]) -> list[dict]:
    return redcap_post(token, {
        "content": "record",
        "format": "json",
        "type": "flat",
        "fields": ",".join(fields),
        "rawOrLabel": "raw",
        "rawOrLabelHeaders": "raw",
        "exportSurveyFields": "false",
        "exportDataAccessGroups": "false",
        "returnFormat": "json",
    }).json()


def choice_map(metadata: list[dict], field_name: str) -> dict[str, str]:
    item = next((m for m in metadata if m.get("field_name") == field_name), None)
    if not item:
        return {}
    raw = item.get("select_choices_or_calculations", "") or ""
    out: dict[str, str] = {}
    for part in raw.split("|"):
        if "," not in part:
            continue
        code, label = part.split(",", 1)
        out[code.strip()] = label.strip()
    return out


def master_rows(rows: list[dict]) -> list[dict]:
    return [r for r in rows if not str(r.get("redcap_repeat_instrument", "")).strip()]


def parse_iso_date(value: str) -> date | None:
    value = str(value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    return None


def fiscal_year_bounds(today: date) -> tuple[date, date, str]:
    start_year = today.year if today.month >= 7 else today.year - 1
    start = date(start_year, 7, 1)
    end = date(start_year + 1, 6, 30)
    return start, end, f"{start_year}/{str(start_year + 1)[-2:]}"


def load_state() -> dict:
    if not STATE_PATH.exists():
        return {"schema_version": 1, "bootstrap_complete": False, "courses": {}}
    data = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    data.setdefault("schema_version", 1)
    data.setdefault("bootstrap_complete", False)
    data.setdefault("courses", {})
    return data


def save_state(state: dict) -> None:
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    tmp = STATE_PATH.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(state, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    tmp.replace(STATE_PATH)


def clean_filename(text: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", text.strip())
    return value.strip("_")[:90] or "course"


def as_yes_no(raw: str) -> str:
    value = str(raw or "").strip()
    if value == "1":
        return "Yes"
    if value == "0":
        return "No"
    return ""


def add_import_comment(ws) -> None:
    ws["A1"].comment = Comment(IMPORT_COMMENT, "MUHAS DCEPD")


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_workbook(course: dict, applicants: list[dict]) -> bytes:
    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "All Applicants"

    all_headers = [
        "record_id", "application_date", "full_name", "email", "phone_number",
        "institution_name", "job_title", "payment_reference", "fee_verified",
        "selected_for_batch", "selected_batch_id",
    ]
    ws_all.append(all_headers)
    for r in applicants:
        ws_all.append([
            r.get("record_id", ""),
            r.get("application_date", ""),
            r.get("full_name", ""),
            r.get("email", ""),
            r.get("phone_number", ""),
            r.get("institution_name", ""),
            r.get("job_title", ""),
            r.get("payment_reference", ""),
            as_yes_no(r.get("fee_verified", "")),
            as_yes_no(r.get("selected_for_batch", "")),
            r.get("selected_batch_id", ""),
        ])
    ws_all.freeze_panes = "A2"
    ws_all.auto_filter.ref = ws_all.dimensions
    set_widths(ws_all, {"A": 14, "B": 15, "C": 28, "D": 28, "E": 18, "F": 28, "G": 24, "H": 24, "I": 15, "J": 20, "K": 28})

    ws_sel = wb.create_sheet("Selection Return")
    sel_headers = [
        "record_id", "full_name", "selected_for_batch", "selected_batch_id",
        "batch_role", "selection_date", "selection_by",
    ]
    ws_sel.append(sel_headers)
    add_import_comment(ws_sel)
    for r in applicants:
        ws_sel.append([
            r.get("record_id", ""),
            r.get("full_name", ""),
            r.get("selected_for_batch", ""),
            r.get("selected_batch_id", ""),
            r.get("batch_role", "") or ("1" if str(r.get("selected_for_batch", "")).strip() == "1" else ""),
            r.get("selection_date", ""),
            r.get("selection_by", ""),
        ])
    ws_sel.freeze_panes = "A2"
    ws_sel.auto_filter.ref = ws_sel.dimensions
    set_widths(ws_sel, {"A": 14, "B": 30, "C": 22, "D": 30, "E": 16, "F": 16, "G": 24})
    dv_yesno = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
    dv_role = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws_sel.add_data_validation(dv_yesno)
    ws_sel.add_data_validation(dv_role)
    if ws_sel.max_row >= 2:
        dv_yesno.add(f"C2:C{ws_sel.max_row}")
        dv_role.add(f"E2:E{ws_sel.max_row}")

    ws_cert = wb.create_sheet("Cert-Graduands Return")
    cert_headers = [
        "record_id", "full_name", "gepg_control_no", "attendance_verified",
        "attendance_verified_date", "certificate_approved", "certification_notes",
    ]
    ws_cert.append(cert_headers)
    add_import_comment(ws_cert)
    selected = [
        r for r in applicants
        if str(r.get("selected_for_batch", "")).strip() == "1"
        and str(r.get("batch_role", "")).strip() in {"", "1"}
    ]
    for r in selected:
        ws_cert.append([
            r.get("record_id", ""),
            r.get("full_name", ""),
            r.get("gepg_control_no", ""),
            r.get("attendance_verified", ""),
            r.get("attendance_verified_date", ""),
            r.get("certificate_approved", ""),
            r.get("certification_notes", ""),
        ])
    ws_cert.freeze_panes = "A2"
    ws_cert.auto_filter.ref = ws_cert.dimensions
    set_widths(ws_cert, {"A": 14, "B": 30, "C": 24, "D": 22, "E": 24, "F": 22, "G": 45})
    dv_att = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
    dv_cert = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
    ws_cert.add_data_validation(dv_att)
    ws_cert.add_data_validation(dv_cert)
    if ws_cert.max_row >= 2:
        dv_att.add(f"D2:D{ws_cert.max_row}")
        dv_cert.add(f"F2:F{ws_cert.max_row}")

    course_note = (
        f"Course: {course.get('course_code','')} — {course.get('course_name','')}\n"
        f"Course Director: {course.get('course_director_name','')}\n"
        f"Vote code: {course.get('vote_code','') or 'Not recorded'}"
    )
    ws_all["A1"].comment = Comment(course_note, "MUHAS DCEPD")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def smtp_settings() -> tuple[str, int, str, str, list[str]]:
    host = os.getenv("DCEPD_SMTP_HOST", "").strip()
    user = os.getenv("DCEPD_SMTP_USERNAME", "").strip()
    password = os.getenv("DCEPD_SMTP_PASSWORD", "").strip()
    port = int(os.getenv("DCEPD_SMTP_PORT", "").strip() or "587")
    cc = []
    for key in ("DCEPD_EMAIL_1", "DCEPD_EMAIL_2", "DCEPD_EMAIL_3", "DCEPD_EMAIL_4"):
        value = os.getenv(key, "").strip()
        if value and value not in cc:
            cc.append(value)
    if not host or not user or not password:
        raise RuntimeError("DCEPD SMTP settings are incomplete.")
    if len(cc) != 4:
        raise RuntimeError("Expected four configured DCEPD CC recipients.")
    return host, port, user, password, cc


def send_pack(course: dict, workbook: bytes, filename: str, scope_label: str, applicant_count: int) -> None:
    host, port, user, password, cc = smtp_settings()
    to_addr = course["contact_email"].strip()
    msg = EmailMessage()
    msg["Subject"] = f"DCEPD Applicant Selection Pack — {course['course_name']} — {scope_label}"
    msg["From"] = user
    msg["To"] = to_addr
    msg["Cc"] = ", ".join(cc)
    msg.set_content(
        f"Dear {course['course_director_name']},\n\n"
        f"Please find attached the MUHAS DCEPD Course Director Applicant Pack for:\n"
        f"{course['course_code']} — {course['course_name']}\n\n"
        f"Applicants included: {applicant_count}\n\n"
        "Please review the All Applicants tab and complete the relevant return tab(s). "
        "Return the Excel workbook to DCEPD. The DCEPD Coordinator/Admin will convert "
        "the relevant return tab to CSV UTF-8 and import it into REDCap.\n\n"
        "Regards,\nMUHAS DCEPD"
    )
    msg.add_attachment(
        workbook,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )
    with smtplib.SMTP(host, port, timeout=60) as smtp:
        smtp.starttls(context=ssl.create_default_context())
        smtp.login(user, password)
        smtp.send_message(msg)


def send_qc_email(qc_csv: bytes, issue_count: int) -> None:
    host, port, user, password, cc = smtp_settings()
    msg = EmailMessage()
    msg["Subject"] = "DCEPD weekly Course Director contact QC"
    msg["From"] = user
    msg["To"] = ", ".join(cc)
    msg.set_content(
        f"The Applicant Pack QC found {issue_count} course/contact exception(s). "
        "No Applicant Pack was sent for affected courses. Please correct Project 75 "
        "director/contact metadata before the next run."
    )
    msg.add_attachment(qc_csv, maintype="text", subtype="csv", filename="applicant_pack_qc.csv")
    with smtplib.SMTP(host, port, timeout=60) as smtp:
        smtp.starttls(context=ssl.create_default_context())
        smtp.login(user, password)
        smtp.send_message(msg)


def write_qc(rows: list[dict]) -> bytes:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    fields = [
        "course_record_id", "course_code", "course_name", "course_director_id",
        "course_director_name", "contact_email", "issue",
    ]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields)
    w.writeheader()
    w.writerows(rows)
    data = buf.getvalue().encode("utf-8-sig")
    (OUT_DIR / "applicant_pack_qc.csv").write_bytes(data)
    return data


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="Build packs/QC but do not send email or update state.")
    args = parser.parse_args()

    today = datetime.now(TZ).date()
    fy_start, fy_end, fy_label = fiscal_year_bounds(today)
    is_friday = today.weekday() == 4

    p75_metadata = export_metadata(P75_TOKEN)
    p75_rows = master_rows(export_records(P75_TOKEN, P75_FIELDS))
    p79_rows = master_rows(export_records(P79_TOKEN, P79_FIELDS))
    director_labels = choice_map(p75_metadata, "course_director_id")

    courses: dict[str, dict] = {}
    email_directors: dict[str, set[str]] = defaultdict(set)
    for r in p75_rows:
        rid = str(r.get("record_id", "")).strip()
        if not rid:
            continue
        code = str(r.get("course_director_id", "")).strip()
        label = director_labels.get(code, code)
        course = dict(r)
        course["course_director_name"] = label
        courses[rid] = course
        email = str(r.get("contact_email", "")).strip().lower()
        if email and code:
            email_directors[email].add(code)

    apps_by_course: dict[str, list[dict]] = defaultdict(list)
    for r in p79_rows:
        course_id = str(r.get("applied_course_id", "")).strip()
        if course_id:
            apps_by_course[course_id].append(r)

    state = load_state()
    bootstrap = not bool(state.get("bootstrap_complete"))
    qc: list[dict] = []
    sent_count = 0
    eligible_count = 0

    for course_id in sorted(apps_by_course, key=lambda x: int(x) if x.isdigit() else x):
        all_apps = sorted(apps_by_course[course_id], key=lambda r: (str(r.get("application_date", "")), str(r.get("record_id", ""))))
        course = courses.get(course_id)
        if not course:
            qc.append({
                "course_record_id": course_id, "course_code": "", "course_name": "",
                "course_director_id": "", "course_director_name": "", "contact_email": "",
                "issue": "Project 79 applied_course_id has no matching Project 75 master course.",
            })
            continue

        director_code = str(course.get("course_director_id", "")).strip()
        director_name = str(course.get("course_director_name", "")).strip()
        email = str(course.get("contact_email", "")).strip()

        issues = []
        if not director_code or any(x in director_name.lower() for x in PLACEHOLDER_DIRECTOR_TEXT):
            issues.append("Course director not identified.")
        if not email:
            issues.append("Course director contact email missing.")
        elif not EMAIL_RE.match(email):
            issues.append("Course director contact email is invalid.")
        elif len(email_directors[email.lower()]) > 1:
            issues.append("Contact email is ambiguous across more than one director code.")
        if not str(course.get("course_code", "")).strip():
            issues.append("Course code missing.")

        if issues:
            for issue in issues:
                qc.append({
                    "course_record_id": course_id,
                    "course_code": course.get("course_code", ""),
                    "course_name": course.get("course_name", ""),
                    "course_director_id": director_code,
                    "course_director_name": director_name,
                    "contact_email": email,
                    "issue": issue,
                })
            continue

        if bootstrap:
            scope_apps = all_apps
            scope_label = "Initial Applicant Pack"
        else:
            scope_apps = [
                r for r in all_apps
                if (d := parse_iso_date(r.get("application_date", ""))) is not None
                and fy_start <= d <= fy_end
            ]
            scope_label = f"FY {fy_label}"

        if not scope_apps:
            continue

        course_state = state["courses"].setdefault(course_id, {"sent_record_ids": [], "last_sent_at": ""})
        sent_ids = set(str(x) for x in course_state.get("sent_record_ids", []))
        scope_ids = [str(r.get("record_id", "")).strip() for r in scope_apps if str(r.get("record_id", "")).strip()]
        new_ids = [rid for rid in scope_ids if rid not in sent_ids]

        if bootstrap:
            should_send = bool(new_ids)
        else:
            should_send = len(new_ids) >= 5 or (is_friday and 1 <= len(new_ids) <= 4)

        if not should_send:
            continue

        eligible_count += 1
        filename = f"DCEPD_Applicant_Pack_{clean_filename(course['course_code'])}_{today.isoformat()}.xlsx"
        workbook = build_workbook(course, scope_apps)
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        (OUT_DIR / filename).write_bytes(workbook)

        if args.dry_run:
            print(f"DRY_RUN course={course_id} applicants={len(scope_apps)} new={len(new_ids)} file={filename}")
            continue

        send_pack(course, workbook, filename, scope_label, len(scope_apps))
        course_state["sent_record_ids"] = sorted(set(sent_ids).union(scope_ids))
        course_state["last_sent_at"] = datetime.now(TZ).isoformat(timespec="seconds")
        course_state["last_scope"] = scope_label
        sent_count += 1
        print(f"SENT course={course_id} applicants={len(scope_apps)} new={len(new_ids)} to={email}")

    qc_csv = write_qc(qc)

    if not args.dry_run:
        if bootstrap and not qc:
            fully_covered = True
            for course_id, rows in apps_by_course.items():
                if course_id not in courses:
                    fully_covered = False
                    break
                sent_ids = set(str(x) for x in state["courses"].get(course_id, {}).get("sent_record_ids", []))
                row_ids = {str(r.get("record_id", "")).strip() for r in rows if str(r.get("record_id", "")).strip()}
                if not row_ids.issubset(sent_ids):
                    fully_covered = False
                    break
            state["bootstrap_complete"] = fully_covered

        state["last_run_at"] = datetime.now(TZ).isoformat(timespec="seconds")
        state["last_run_qc_issues"] = len(qc)
        save_state(state)

        if is_friday and qc:
            send_qc_email(qc_csv, len(qc))

    print(json.dumps({
        "date": today.isoformat(),
        "bootstrap": bootstrap,
        "eligible_packs": eligible_count,
        "sent_packs": sent_count,
        "qc_issues": len(qc),
        "state_path": str(STATE_PATH),
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
